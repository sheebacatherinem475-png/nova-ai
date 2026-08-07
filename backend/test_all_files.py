import httpx
import json
import time
import os
import pandas as pd
from fpdf import FPDF
import docx

URL = 'https://nova-ai-0m0x.onrender.com'

def create_files():
    # TXT
    with open('test.txt', 'w') as f:
        f.write('This document is about Nova AI. It is a TXT file.')
    
    # PDF
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.cell(200, 10, txt="This document is about Nova AI. It is a PDF file.", ln=True, align='L')
    pdf.output("test.pdf")

    # DOCX
    doc = docx.Document()
    doc.add_paragraph('This document is about Nova AI. It is a DOCX file.')
    doc.save('test.docx')
    
    # JSON
    with open('test.json', 'w') as f:
        json.dump([{"id": 1, "data": "This document is about Nova AI. It is a JSON file."}], f)
        
    # CSV
    df = pd.DataFrame({"id": [1], "data": ["This document is about Nova AI. It is a CSV file."]})
    df.to_csv('test.csv', index=False)
    
    # XLSX
    df.to_excel('test.xlsx', index=False)

create_files()

def run_test():
    try:
        print('Checking if backend is up and deployed...')
        while True:
            try:
                res = httpx.get(f'{URL}/health', timeout=10)
                if res.status_code == 200:
                    break
            except Exception:
                pass
            time.sleep(5)
            
        print('Registering user...')
        httpx.post(f'{URL}/api/auth/register', json={'email': 'livedebug_all@test.com', 'password': 'password'}, timeout=60.0)
        
        res = httpx.post(f'{URL}/api/auth/login', data={'username': 'livedebug_all@test.com', 'password': 'password'}, timeout=60.0)
        token = res.json().get('access_token')
        if not token:
            print('Login failed!', res.text)
            return
            
        headers = {'Authorization': f'Bearer {token}'}
            
        doc_files = [
            ('test.txt', 'text/plain'),
            ('test.pdf', 'application/pdf'),
            ('test.docx', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')
        ]
        
        for filename, mime in doc_files:
            print(f'\n--- Testing Document Upload: {filename} ---')
            with open(filename, 'rb') as f:
                upload_files = {'file': (filename, f, mime)}
                res = httpx.post(f'{URL}/api/documents/upload', files=upload_files, headers=headers, timeout=120.0)
                print(f'Upload response ({res.status_code}):', res.text)
        
        dataset_files = [
            ('test.csv', 'text/csv'),
            ('test.json', 'application/json'),
            ('test.xlsx', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        ]
        
        uploaded_datasets = []
        for filename, mime in dataset_files:
            print(f'\n--- Testing Dataset Upload: {filename} ---')
            with open(filename, 'rb') as f:
                upload_files = {'file': (filename, f, mime)}
                res = httpx.post(f'{URL}/api/datasets/upload', files=upload_files, headers=headers, timeout=120.0)
                print(f'Upload response ({res.status_code}):', res.text)
                if res.status_code == 200:
                    uploaded_datasets.append(res.json().get('id'))
        
        print('\n--- Testing Chat RAG (Documents) ---')
        payload = {'message': 'What are the uploaded documents about?', 'history': []}
        chat_headers = {'Authorization': f'Bearer {token}', 'Content-Type': 'application/json'}
        
        with httpx.stream('POST', f'{URL}/api/chat', json=payload, headers=chat_headers, timeout=120.0) as r:
            for chunk in r.iter_text():
                print(chunk, end='')
        print()
        
        if uploaded_datasets:
            dataset_id = uploaded_datasets[0]
            print(f'\n--- Testing Data Analysis (Dataset {dataset_id}) ---')
            payload = {'message': 'What is in this dataset?', 'history': [], 'dataset_id': dataset_id}
            
            with httpx.stream('POST', f'{URL}/api/data-analysis/analyze', json=payload, headers=chat_headers, timeout=120.0) as r:
                for chunk in r.iter_text():
                    print(chunk, end='')
            print()
        
    except Exception as e:
        print('Error:', e)

run_test()
